import docx

doc = docx.Document('d:/Documents/teaching/OBE2-CS69-3.docx')
with open('d:/Documents/curriculum69/scripts/ge_info.txt', 'w', encoding='utf-8') as f:
    for i, p in enumerate(doc.paragraphs):
        text = p.text.strip()
        if any(k in text for k in ['ศึกษาทั่วไป', '001101', '001102', '001225', 'B1', 'CEFR', 'ภาษาและการสื่อสาร', 'กลุ่มวิชาภาษา']):
            f.write(f"P{i}: {text}\n")
            # print surrounding paragraphs
            for j in range(max(0, i-2), min(len(doc.paragraphs), i+8)):
                f.write(f"  [{j}] {doc.paragraphs[j].text.strip()}\n")
            f.write("-" * 50 + "\n")
            
    for ti, table in enumerate(doc.tables):
        found = False
        for row in table.rows:
            row_text = ' | '.join(cell.text.strip().replace('\n', ' ') for cell in row.cells)
            if any(k in row_text for k in ['ศึกษาทั่วไป', '001101', '001225', 'กลุ่มวิชาภาษา']):
                found = True
                break
        if found:
            f.write(f"\n=== TABLE {ti} ===\n")
            for row in table.rows:
                f.write(' | '.join(cell.text.strip().replace('\n', ' ') for cell in row.cells) + "\n")
print("Done extracting GE info")
